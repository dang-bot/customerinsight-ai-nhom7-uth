import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report_docx(md_path, docx_path):
    doc = docx.Document()
    
    # Page setup - Standard A4 with 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(4)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []
    
    for line in lines:
        raw = line.strip()
        
        # Check table
        if raw.startswith('|') and raw.endswith('|'):
            in_table = True
            table_lines.append(raw)
            continue
        elif in_table:
            # Process table
            render_table(doc, table_lines)
            in_table = False
            table_lines = []

        if not raw:
            continue

        # Horizontal rule
        if raw == '---':
            continue

        # Title and Headers
        if raw.startswith('# '):
            text = raw[2:].replace('**', '').replace('*', '')
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            
        elif raw.startswith('## '):
            text = raw[3:].replace('**', '').replace('*', '')
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy Blue
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            
        elif raw.startswith('### '):
            text = raw[4:].replace('**', '').replace('*', '')
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            
        elif raw.startswith('#### '):
            text = raw[5:].replace('**', '').replace('*', '')
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            
        elif raw.startswith('> '):
            text = raw[2:].replace('**', '').replace('*', '')
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.italic = True
            run.font.color.rgb = RGBColor(100, 116, 139)
            
        elif raw.startswith('* ') or raw.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            format_inline(p, raw[2:])
            
        elif re.match(r'^\d+\.\s', raw):
            match = re.match(r'^\d+\.\s', raw)
            p = doc.add_paragraph(style='List Number')
            format_inline(p, raw[match.end():])
            
        elif raw.startswith('```'):
            continue
        else:
            p = doc.add_paragraph()
            format_inline(p, raw)

    if in_table and table_lines:
        render_table(doc, table_lines)

    doc.save(docx_path)
    print("Successfully generated Word document!")

def format_inline(p, text):
    # Splits by bold tags **text**
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.font.bold = True
        else:
            # Inline code or math cleanups
            clean_part = part.replace('`', '').replace('$', '')
            p.add_run(clean_part)

def render_table(doc, lines):
    # Filter header separator line like |:---:|:---|
    rows_data = []
    for line in lines:
        if re.search(r'\|[:\-\s]+\|', line):
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        rows_data.append(cells)
        
    if not rows_data:
        return
        
    num_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    
    for row_idx, row in enumerate(rows_data):
        for col_idx, text in enumerate(row):
            if col_idx < num_cols:
                cell = table.cell(row_idx, col_idx)
                cell.text = text.replace('**', '')
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                # Header row styling
                if row_idx == 0:
                    set_cell_background(cell, "1E3A8A") # Navy Blue
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.name = 'Times New Roman'
                            run.font.bold = True
                            run.font.size = Pt(11)
                            run.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                    set_cell_background(cell, bg_color)
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(11)
                            run.font.color.rgb = RGBColor(30, 41, 59)
                            
    doc.add_paragraph() # Spacing

if __name__ == "__main__":
    md_file = r"D:\UTH\Nam hai(2025-2026)\HK_he\TriTueNhanTao\Thực hành\customerinsight-ai\docs\BaoCao_Nhom7_CustomerInsightAI.md"
    doc_file = r"D:\UTH\Nam hai(2025-2026)\HK_he\TriTueNhanTao\Thực hành\customerinsight-ai\docs\BaoCao_Nhom7_CustomerInsightAI.docx"
    create_report_docx(md_file, doc_file)
